"""
Flujo combinado: Extracción de PDFs y generación directa de Excel
"""

import os
import sys
import time
from typing import List, Dict, Optional
from utils.logger import Logger
# Agregar el directorio padre al path para importar módulos
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Importar extractores específicos
from extractores.extractor_afp import extraer_datos_afp
from extractores.extractor_boleta import extraer_datos_boleta
from extractores.extractor_quinta import extraer_datos_quinta
logger = Logger("CorePipeline3_Diagnostico")
# ============================================
# CONFIGURACIÓN DE CARPETAS Y TIPOS
# ============================================

CARPETAS_CONFIG = {
    "1_Boletas": {
        "tipo": "BOLETA",
        "extractor": extraer_datos_boleta,
        "campos_extra": ["fecha_extraida"]
    },
    "2_Afp": {
        "tipo": "AFP",
        "extractor": extraer_datos_afp,
        "campos_extra": []
    },
    "3_5ta": {
        "tipo": "QUINTA",
        "extractor": extraer_datos_quinta,
        "campos_extra": []
    }
}

# ============================================
# FUNCIONES DE PROCESAMIENTO
# ============================================

def procesar_carpeta(ruta_carpeta: str, config: Dict) -> List[Dict]:
    """
    Procesa todos los PDFs en una carpeta y genera registros.
    Progreso mostrado por lotes de 100 archivos.
    """
    inicio = time.time()
    registros = []
    tipo_documento = config["tipo"]
    extractor = config["extractor"]

    archivos_pdf = [f for f in os.listdir(ruta_carpeta) if f.lower().endswith('.pdf')]

    print(f"\n📂 Procesando carpeta: {os.path.basename(ruta_carpeta)}")
    logger.info(f"📂 Procesando carpeta: {os.path.basename(ruta_carpeta)}")
    print(f"   Tipo de documento: {tipo_documento}")
    logger.info(f"   Tipo: {tipo_documento}")
    print(f"   Total de PDFs encontrados: {len(archivos_pdf)}")
    logger.info(f"   PDFs: {len(archivos_pdf)}")

    for idx, archivo in enumerate(archivos_pdf, 1):
        ruta_completa = os.path.join(ruta_carpeta, archivo)
        resultado = extractor(ruta_completa)

        registro = {
            "archivo_original": archivo,
            "tipo_documento": tipo_documento,
            "nombre_extraido": resultado.get("nombre"),
            "dni_extraido": resultado.get("dni"),
            "exito_extraccion": resultado.get("exito", False),
            "observaciones": resultado.get("observaciones", "")
        }

        if "fecha" in resultado and resultado["fecha"]:
            registro["fecha_extraida"] = resultado["fecha"]

        registros.append(registro)

        # Mostrar progreso cada 100 archivos
        if idx % 100 == 0:
            print(f"   Procesados: {idx}/{len(archivos_pdf)}")
            logger.info(f"   Procesados: {idx}/{len(archivos_pdf)}")

    fin = time.time()
    tiempo_total = fin - inicio
    print(f"   ✅ Completado: {len(registros)} registros en {int(tiempo_total // 60)}m {int(tiempo_total % 60)}s")
    logger.info(f"   ✅ Completado: {len(registros)} registros en {int(tiempo_total // 60)}m {int(tiempo_total % 60)}s")

    return registros

# ============================================
# FUNCIONES DE EXCEL
# ============================================

def generar_excel_multihoja(datos_por_hoja: Dict[str, List[Dict]], ruta_excel: str) -> bool:
    """
    Genera Excel con múltiples hojas directamente desde los datos extraídos.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        for nombre_hoja, registros in datos_por_hoja.items():
            ws = wb.create_sheet(title=nombre_hoja)
            if not registros:
                continue

            encabezados = list(registros[0].keys())
            for col_idx, encabezado in enumerate(encabezados, 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.value = encabezado.replace("_", " ").upper()
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            for row_idx, registro in enumerate(registros, 2):
                for col_idx, encabezado in enumerate(encabezados, 1):
                    valor = registro.get(encabezado, "")
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.value = valor

                    if encabezado in ["tipo_documento", "exito_extraccion", "dni_extraido"]:
                        cell.alignment = Alignment(horizontal="center")

                    if encabezado == "exito_extraccion":
                        if valor:
                            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                        else:
                            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

            for col_idx, encabezado in enumerate(encabezados, 1):
                column_letter = openpyxl.utils.get_column_letter(col_idx)
                if encabezado == "archivo_original":
                    ws.column_dimensions[column_letter].width = 40
                elif encabezado == "observaciones":
                    ws.column_dimensions[column_letter].width = 35
                elif encabezado == "nombre_extraido":
                    ws.column_dimensions[column_letter].width = 30
                elif encabezado == "tipo_documento":
                    ws.column_dimensions[column_letter].width = 15
                elif encabezado == "dni_extraido":
                    ws.column_dimensions[column_letter].width = 12
                elif encabezado == "fecha_extraida":
                    ws.column_dimensions[column_letter].width = 15
                elif encabezado == "exito_extraccion":
                    ws.column_dimensions[column_letter].width = 18
                else:
                    ws.column_dimensions[column_letter].width = 20

            ws.freeze_panes = "A2"

        wb.save(ruta_excel)
        print(f"   ✅ Excel generado: {ruta_excel}")
        logger.info(f"   ✅ Excel generado: {ruta_excel}")
        return True

    except Exception as e:
        print(f"   ❌ Error al generar Excel: {e}")
        logger.error(f"   ❌ Error al generar Excel: {e}")
        import traceback
        traceback.print_exc()
        logger.error(traceback.format_exc())
        return False

# ============================================
# FUNCIÓN PRINCIPAL
# ============================================

def procesar_diagnostico_a_excel(ruta_carpeta_trabajo: str, ruta_excel_final: str, guardar_json_opcional: bool = False):
    """
    Procesa PDFs y genera Excel directamente. 
    Opción de guardar JSON para debugging.
    """
    print("="*60)
    print("🚀 PROCESO DE DIAGNÓSTICO DIRECTO A EXCEL")
    print("="*60)
    logger.info("="*60)
    logger.info("🚀 PROCESO DE DIAGNÓSTICO DIRECTO A EXCEL")
    logger.info("="*60)

    if not os.path.isdir(ruta_carpeta_trabajo):
        print(f"❌ Carpeta '{ruta_carpeta_trabajo}' no existe")
        logger.error(f"❌ Carpeta '{ruta_carpeta_trabajo}' no existe")
        return

    datos_por_hoja = {}

    for nombre_carpeta, config in CARPETAS_CONFIG.items():
        ruta_subcarpeta = os.path.join(ruta_carpeta_trabajo, nombre_carpeta)
        if not os.path.isdir(ruta_subcarpeta):
            print(f"⚠️ Carpeta '{nombre_carpeta}' no encontrada, omitiendo...")
            logger.warning(f"⚠️ Carpeta '{nombre_carpeta}' no encontrada, omitiendo...")
            continue

        registros = procesar_carpeta(ruta_subcarpeta, config)
        datos_por_hoja[nombre_carpeta] = registros

        if guardar_json_opcional:
            import json
            ruta_json = os.path.join(ruta_carpeta_trabajo, f"diagnostico_{nombre_carpeta}.json")
            with open(ruta_json, 'w', encoding='utf-8') as f:
                json.dump(registros, f, ensure_ascii=False, indent=2)
            print(f"   💾 JSON opcional guardado: {ruta_json}")
            logger.info(f"   💾 JSON opcional guardado: {ruta_json}")

    if datos_por_hoja:
        generar_excel_multihoja(datos_por_hoja, ruta_excel_final)

# ============================================
# OBTENER RUTA DE CARPETA
# ============================================

def obtener_ruta_carpeta(ruta: Optional[str] = None) -> str:
    if ruta and os.path.isdir(ruta):
        return os.path.normpath(ruta)
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        root.attributes('-topmost', True)
        ruta_seleccionada = filedialog.askdirectory(title="Seleccionar carpeta de trabajo", mustexist=True)
        root.destroy()
        return os.path.normpath(ruta_seleccionada) if ruta_seleccionada else None
    except Exception as e:
        print(f"❌ Error al abrir explorador: {e}")
        return None

# ============================================
# EJECUCIÓN PRINCIPAL
# ============================================

if __name__ == "__main__":
    ruta_trabajo = obtener_ruta_carpeta()
    if ruta_trabajo:
        # Generar timestamp para el nombre del archivo
        timestamp = time.strftime("%d.%m.%Y_%H.%M.%S")
        nombre_excel = f"diagnostico_consolidado_{timestamp}.xlsx"
        ruta_excel = os.path.join(ruta_trabajo, nombre_excel)
        procesar_diagnostico_a_excel(ruta_trabajo, ruta_excel, guardar_json_opcional=False)
    else:
        print("⛔ No se seleccionó carpeta. Operación cancelada.")